"""Parser behaviour around multiple roles at one employer (promotion /
internal transfer). The prose rule has been in PARSE_SYSTEM for a while but
resumes still came back with the earlier role dropped — a concrete worked
example in the prompt makes the fast model follow it far more reliably.
"""
import json

import pytest
from unittest.mock import AsyncMock, MagicMock

from app.services.resume_parser import extract_text_from_docx, extract_text_from_pdf, parse_resume_text, PARSE_SYSTEM


def _provider_returning(payload: dict) -> MagicMock:
    provider = MagicMock()
    provider.complete = AsyncMock(return_value=json.dumps(payload))
    return provider


def test_parse_system_prompt_carries_a_worked_multi_role_example():
    # A named company that appears at least twice inside a single worked
    # example block — proof the prompt shows (not just tells) how a promotion
    # becomes two separate entries.
    assert "Example" in PARSE_SYSTEM
    assert PARSE_SYSTEM.count("Northwind") >= 2


def test_parse_system_prompt_tells_the_model_to_use_hyperlink_urls():
    assert "Hyperlinks found in this document" in PARSE_SYSTEM


# ---------------------------------------------------------------------------
# extract_text_from_pdf: a PDF hyperlink's visible anchor text (e.g. the bare
# word "GitHub") never contains the actual destination URL — pypdf's
# extract_text() only returns what's visibly printed. Without pulling the
# URI from the link annotation too, the parser prompt above never sees the
# real github.com/linkedin.com URL, only the word "GitHub"/"LinkedIn", so it
# has nothing to put in contact.github/contact.linkedin. See _extract_link_uris.
# ---------------------------------------------------------------------------

weasyprint = pytest.importorskip("weasyprint")


def _pdf_with_link(anchor_text: str, href: str) -> bytes:
    import weasyprint as wp
    html = f'<html><body><p>Contact: <a href="{href}">{anchor_text}</a></p></body></html>'
    return wp.HTML(string=html).write_pdf()


def test_extract_text_from_pdf_surfaces_hyperlink_url_behind_bare_anchor_text():
    pdf_bytes = _pdf_with_link("GitHub", "https://github.com/janedoe/my-repo")
    text = extract_text_from_pdf(pdf_bytes)
    assert "GitHub" in text
    assert "https://github.com/janedoe/my-repo" in text


def test_extract_text_from_pdf_dedupes_repeated_hyperlinks():
    import weasyprint as wp
    html = (
        '<html><body>'
        '<p><a href="https://github.com/janedoe/repo-one">GitHub</a></p>'
        '<p><a href="https://github.com/janedoe/repo-one">GitHub</a></p>'
        '</body></html>'
    )
    pdf_bytes = wp.HTML(string=html).write_pdf()
    text = extract_text_from_pdf(pdf_bytes)
    assert text.count("https://github.com/janedoe/repo-one") == 1


def test_extract_text_from_pdf_omits_hyperlink_section_when_no_links():
    import weasyprint as wp
    pdf_bytes = wp.HTML(string="<html><body><p>Jane Doe, no links here.</p></body></html>").write_pdf()
    text = extract_text_from_pdf(pdf_bytes)
    assert "Hyperlinks found" not in text


# ---------------------------------------------------------------------------
# Same bug, same fix, for Word resumes: a hyperlink's visible run text (e.g.
# "GitHub") never includes the address it points at — see the DOCX branch of
# extract_text_from_docx / _append_link_hints.
# ---------------------------------------------------------------------------


def _add_hyperlink(paragraph, url: str, text: str) -> None:
    """python-docx (this version) has no writer-side add_hyperlink helper —
    build the run manually via its relationship + oxml APIs, purely for
    constructing a realistic test fixture."""
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    run.append(OxmlElement("w:rPr"))
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _docx_with_link(anchor_text: str, href: str) -> bytes:
    import io as _io
    import docx

    doc = docx.Document()
    p = doc.add_paragraph("Contact: ")
    _add_hyperlink(p, href, anchor_text)
    buf = _io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_extract_text_from_docx_surfaces_hyperlink_url_behind_bare_anchor_text():
    docx_bytes = _docx_with_link("GitHub", "https://github.com/janedoe/my-repo")
    text = extract_text_from_docx(docx_bytes)
    assert "GitHub" in text
    assert "https://github.com/janedoe/my-repo" in text


def test_extract_text_from_docx_omits_hyperlink_section_when_no_links():
    import io as _io
    import docx

    doc = docx.Document()
    doc.add_paragraph("Jane Doe, no links here.")
    buf = _io.BytesIO()
    doc.save(buf)
    text = extract_text_from_docx(buf.getvalue())
    assert "Hyperlinks found" not in text


@pytest.mark.asyncio
async def test_two_roles_at_one_company_survive_as_separate_entries():
    payload = {
        "contact": {"name": "Jane Doe", "email": "jane@example.com"},
        "experience": [
            {
                "company": "Acme Corp", "title": "Senior Software Engineer",
                "start": "Jan 2023", "end": "Present",
                "bullets": ["Led the payments platform rewrite."],
            },
            {
                "company": "Acme Corp", "title": "Software Engineer",
                "start": "Jun 2020", "end": "Dec 2022",
                "bullets": ["Built the original checkout service."],
            },
        ],
        "education": [],
        "skills": ["Python"],
    }
    result = await parse_resume_text("<resume text>", _provider_returning(payload))

    exp = result["experience"]
    assert len(exp) == 2
    assert [e["title"] for e in exp] == ["Senior Software Engineer", "Software Engineer"]
    assert [e["company"] for e in exp] == ["Acme Corp", "Acme Corp"]
    # Each role keeps only its own bullets — no cross-contamination or merge.
    assert exp[0]["bullets"] == ["Led the payments platform rewrite."]
    assert exp[1]["bullets"] == ["Built the original checkout service."]
